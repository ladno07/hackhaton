import io
from datetime import date, timedelta
from openpyxl import load_workbook
from sqlalchemy import select
from app.models import Athlete, AthleteCoach, StageTransfer, TrainingStandard, User
from app.services.standards import STAGES
from tests.test_ministry import READ_ONLY, ministry_login

PUSH = "Сгибание и разгибание рук в упоре лёжа"
CONTEST = "Решение задач тренировочного контеста за 5 часов"


def coached(factory):
    """Coached demo athletes in seed order: ready for УТ, halfway to УТ."""
    with factory() as db:
        ids = db.scalars(
            select(AthleteCoach.athlete_id)
            .where(AthleteCoach.status == "active")
            .order_by(AthleteCoach.athlete_id)
        ).all()
        return ids[0], ids[1]


def item(state, name):
    return next(i for i in state["items"] if i["name"] == name)


def test_readiness_compares_results_with_thresholds(env):
    client, factory, headers, users = env
    me = client.get("/api/attestation", headers=headers["athlete"]).json()
    assert len(me) == 1
    mine = client.get(f"/api/attestation/{me[0]['athlete_id']}", headers=headers["athlete"]).json()
    # УТ → ССМ: four of five female standards passed, КМС is missing
    assert mine["stage"] == "ut" and mine["target"] == "ssm" and mine["sex"] == "female"
    assert (mine["passed"], mine["total"]) == (4, 5)
    assert not mine["ready"] and not mine["rank_ok"] and mine["rank_required"] == "КМС"
    push = item(mine, PUSH)
    # the best of two attempts counts, female threshold for ССМ is 10
    assert push["threshold"] == 10 and push["best_value"] == 9 and push["attempts"] == 2
    assert not push["met"] and not mine["can_record"] and not mine["can_transfer"]
    assert all(i["sex"] in (None, "female") for i in mine["items"])

    listed = client.get("/api/attestation", headers=headers["coach"]).json()
    ready_id, halfway_id = coached(factory)
    assert [i["athlete_id"] for i in listed if i["ready"]] == [ready_id]
    halfway = client.get(f"/api/attestation/{halfway_id}", headers=headers["coach"]).json()
    assert halfway["target"] == "ut" and (halfway["passed"], halfway["total"]) == (3, 5)
    assert halfway["can_record"] and not halfway["can_transfer"]
    # the coach sees only linked athletes
    assert client.get(f"/api/attestation/{me[0]['athlete_id']}", headers=headers["coach"]).status_code == 404

    # attention block of the federation lists the ready athlete
    attention = client.get("/api/dashboard/attention", headers=headers["admin"]).json()
    ready = [i for i in attention["items"] if i["kind"] == "stage_transfers"]
    assert [i["id"] for i in ready] == [ready_id] and not ready[0]["overdue"]
    assert attention["counts"]["stage_transfers"] == 1


def test_coach_records_results_and_federation_transfers(env):
    client, factory, headers, users = env
    _, athlete = coached(factory)
    state = client.get(f"/api/attestation/{athlete}", headers=headers["coach"]).json()
    push, contest = item(state, PUSH), item(state, CONTEST)
    url = f"/api/attestation/{athlete}/results"
    today = date.today().isoformat()

    # a male threshold does not apply to a female athlete
    with factory() as db:
        male = db.scalar(
            select(TrainingStandard.id).where(
                TrainingStandard.stage == "ut",
                TrainingStandard.name == PUSH,
                TrainingStandard.sex == "male",
            )
        )
    res = client.post(url, json={"standard_id": male, "value": 30, "tested_on": today}, headers=headers["coach"])
    assert res.status_code == 422
    future = (date.today() + timedelta(days=1)).isoformat()
    res = client.post(url, json={"standard_id": push["id"], "value": 9, "tested_on": future}, headers=headers["coach"])
    assert res.status_code == 422
    assert client.post(url, json={"standard_id": push["id"], "value": 9, "tested_on": today}, headers=headers["athlete"]).status_code == 403

    res = client.post(url, json={"standard_id": push["id"], "value": 9, "tested_on": today}, headers=headers["coach"])
    assert res.status_code == 201, res.text
    assert res.json()["passed"] and res.json()["threshold"] == push["threshold"]
    transfer = {"to_stage": "ut", "decided_on": today, "order_number": "Т-1"}
    res = client.post(f"/api/attestation/{athlete}/transfer", json=transfer, headers=headers["admin"])
    assert res.status_code == 409 and "Не выполнено нормативов: 1" in res.json()["detail"]

    res = client.post(url, json={"standard_id": contest["id"], "value": 5, "tested_on": today}, headers=headers["coach"])
    assert res.status_code == 201
    state = client.get(f"/api/attestation/{athlete}", headers=headers["admin"]).json()
    assert state["ready"] and state["can_transfer"]
    notices = client.get("/api/notifications", headers=headers["coach"]).json()
    assert any(n["event_key"] == f"stage-ready:{athlete}:ut" for n in notices)

    assert client.post(f"/api/attestation/{athlete}/transfer", json={**transfer, "to_stage": "ssm"}, headers=headers["admin"]).status_code == 409
    assert client.post(f"/api/attestation/{athlete}/transfer", json=transfer, headers=headers["coach"]).status_code == 403
    res = client.post(f"/api/attestation/{athlete}/transfer", json=transfer, headers=headers["admin"])
    assert res.status_code == 201, res.text
    done = res.json()
    assert done["from_stage"] == STAGES["np"] and done["to_stage_name"] == STAGES["ut"]
    assert {s["name"] for s in done["basis"]["standards"]} >= {PUSH, CONTEST}

    after = client.get(f"/api/attestation/{athlete}", headers=headers["coach"]).json()
    assert after["stage"] == "ut" and after["target"] == "ssm" and not after["ready"]
    assert [t["order_number"] for t in after["transfers"]] == ["Т-1"]
    # earlier attempts stay in the history with their verdict
    assert {r["standard"] for r in after["results"]} >= {PUSH, CONTEST}
    assert any(not r["passed"] for r in after["results"])
    notices = client.get("/api/notifications", headers=headers["coach"]).json()
    assert any(n["event_key"] == f"stage-transfer:{done['id']}" for n in notices)


def test_athlete_cannot_change_own_stage(env):
    client, factory, headers, users = env
    aid = users["athlete"]["athlete_id"]
    profile = client.get(f"/api/athletes/{aid}", headers=headers["athlete"]).json()
    body = {
        k: profile[k]
        for k in ["full_name", "organization_id", "municipality", "training_stage"]
    }
    body["disciplines"] = [
        {k: d[k] for k in ["discipline_id", "start_year", "end_year"]}
        for d in profile["disciplines"]
    ]
    assert client.put(f"/api/athletes/{aid}", json=body, headers=headers["athlete"]).status_code == 200
    jump = {**body, "training_stage": STAGES["vsm"]}
    res = client.put(f"/api/athletes/{aid}", json=jump, headers=headers["athlete"])
    assert res.status_code == 403
    assert client.put(f"/api/athletes/{aid}", json=jump, headers=headers["admin"]).status_code == 200


def test_threshold_change_and_unknown_sex(env):
    client, factory, headers, users = env
    ready_id, _ = coached(factory)
    with factory() as db:
        standard = db.scalar(
            select(TrainingStandard).where(
                TrainingStandard.stage == "ut",
                TrainingStandard.name == PUSH,
                TrainingStandard.sex == "male",
            )
        )
        body = {
            k: getattr(standard, k)
            for k in ("stage", "category", "name", "unit", "better", "sex", "discipline_id", "source", "sort_order")
        }
    # stricter threshold: the ready athlete (15 push-ups) is no longer ready
    res = client.put(f"/api/standards/{standard.id}", json={**body, "threshold": 16}, headers=headers["admin"])
    assert res.status_code == 200, res.text
    state = client.get(f"/api/attestation/{ready_id}", headers=headers["admin"]).json()
    assert not state["ready"] and not item(state, PUSH)["met"]
    # the recorded attempt keeps the threshold of its test date
    assert next(r for r in state["results"] if r["standard"] == PUSH)["threshold"] == 13
    res = client.put(f"/api/standards/{standard.id}", json={**body, "threshold": 16, "unit": "повт."}, headers=headers["admin"])
    assert res.status_code == 409
    client.put(f"/api/standards/{standard.id}", json={**body, "threshold": 13}, headers=headers["admin"])

    with factory() as db:
        user = db.get(User, db.get(Athlete, ready_id).user_id)
        user.sex = None
        db.commit()
    state = client.get(f"/api/attestation/{ready_id}", headers=headers["admin"]).json()
    assert not state["ready"] and any("Не указан пол" in b for b in state["blockers"])
    # without sex the standards of both sexes are listed and none of them counts
    assert {i["sex"] for i in state["items"]} == {None, "male", "female"}

    archived = client.post(f"/api/standards/{standard.id}/archive", headers=headers["admin"]).json()
    assert archived["archived_at"]
    ids = {s["id"] for s in client.get("/api/standards?stage=ut", headers=headers["athlete"]).json()}
    assert standard.id not in ids
    assert standard.id in {s["id"] for s in client.get("/api/standards?archived=true", headers=headers["admin"]).json()}


def test_catalog_access_and_exports(env):
    client, factory, headers, users = env
    ministry = ministry_login(client)
    catalog = client.get("/api/standards", headers=ministry).json()
    assert {s["stage"] for s in catalog} == set(STAGES)
    assert [s["stage"] for s in catalog] == sorted(
        (s["stage"] for s in catalog), key=list(STAGES).index
    )
    female = client.get("/api/standards?stage=np&sex=female", headers=ministry).json()
    assert female and all(s["sex"] in (None, "female") for s in female)
    meta = client.get("/api/standards/meta", headers=headers["coach"]).json()
    assert [s["rank"] for s in meta["stages"]] == [None, None, "КМС", "МС"]
    assert client.get("/api/standards", headers=headers["judge"]).status_code == 403

    body = {
        "stage": "np", "category": "ofp", "name": "Прыжок в длину с места",
        "unit": "см", "better": "higher", "threshold": 110, "sex": "male",
    }
    res = client.post("/api/standards", json=body, headers=ministry)
    assert res.status_code == 403 and res.json()["detail"] == READ_ONLY
    assert client.post("/api/standards", json=body, headers=headers["coach"]).status_code == 403
    res = client.post("/api/standards", json=body, headers=headers["admin"])
    assert res.status_code == 201 and res.json()["threshold_text"] == "≥ 110 см"

    history = client.get("/api/attestation/history", headers=ministry).json()
    with factory() as db:
        transfer = db.scalars(select(StageTransfer)).one()
    assert [t["id"] for t in history["transfers"]] == [transfer.id]
    assert any(r["expired"] for r in history["results"])

    for fmt in ("pdf", "docx", "xlsx"):
        res = client.get(f"/api/exports/attestation.{fmt}", headers=ministry)
        assert res.status_code == 200, res.text
    sheet = load_workbook(io.BytesIO(res.content)).active
    text = " ".join(str(c.value) for r in sheet.iter_rows() for c in r if c.value)
    assert "Ведомость готовности" in text and "Готов к переводу" in text
    assert client.get("/api/exports/attestation.pdf", headers=headers["coach"]).status_code == 403
