import io
from docx import Document
from openpyxl import load_workbook
from sqlalchemy import select
from app.models import (
    Competition,
    CompetitionJudge,
    Judge,
    ProtocolVersion,
)
from tests.test_ministry import ministry_login


def docx_text(content):
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts += [cell.text for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def setup(factory, users):
    """A published competition: demo judge is chief, admin's card is secretary."""
    with factory() as db:
        chief = db.scalar(select(Judge).where(Judge.user_id == users["judge"]["id"]))
        secretary = db.scalar(select(Judge).where(Judge.user_id == users["admin"]["id"]))
        comp = db.scalars(
            select(CompetitionJudge, Competition)
            .join(Competition, Competition.id == CompetitionJudge.competition_id)
            .where(CompetitionJudge.judge_id == chief.id, Competition.state == "published")
        ).first()
        db.add(
            CompetitionJudge(
                competition_id=comp.competition_id,
                judge_id=secretary.id,
                position="Главный секретарь",
            )
        )
        db.commit()
        return comp.competition_id


def test_panel_and_chief_report_documents(env):
    client, factory, headers, users = env
    comp_id = setup(factory, users)
    with factory() as db:
        comp = db.get(Competition, comp_id)
        latest = db.scalars(
            select(ProtocolVersion)
            .where(ProtocolVersion.competition_id == comp_id)
            .order_by(ProtocolVersion.version.desc())
        ).first()
        winners = [e["full_name"] for e in latest.entries if e["place"] <= 3]
        title = comp.title

    for fmt in ("pdf", "docx", "xlsx"):
        for kind in ("panel", "chief-report"):
            res = client.get(f"/api/exports/{kind}/{comp_id}.{fmt}", headers=headers["admin"])
            assert res.status_code == 200, (kind, fmt, res.text)
            assert len(res.content) > 1000

    panel = docx_text(
        client.get(f"/api/exports/panel/{comp_id}.docx", headers=headers["admin"]).content
    )
    assert "СОСТАВ" in panel.upper() and title in panel
    assert "Главный секретарь" in panel and "УТВЕРЖДАЮ" in panel
    sheet = load_workbook(
        io.BytesIO(
            client.get(f"/api/exports/panel/{comp_id}.xlsx", headers=headers["admin"]).content
        )
    ).active
    positions = [
        row[1]
        for row in sheet.iter_rows(values_only=True)
        if row[1] in ("Главный судья", "Главный секретарь")
    ]
    assert positions == ["Главный судья", "Главный секретарь"]  # chief heads the list

    report = docx_text(
        client.get(f"/api/exports/chief-report/{comp_id}.docx", headers=headers["admin"]).content
    )
    assert "главного судьи" in report and "Призёры соревнования" in report
    assert all(name.removeprefix("Демо · ").split()[0] in report for name in winners)
    assert "Участников по протоколу" in report and "Исправлений протокола" in report


def test_judging_documents_access(env):
    client, factory, headers, users = env
    comp_id = setup(factory, users)
    as_judge_admin = {**headers["admin"], "X-Active-Role": "judge"}
    ministry = ministry_login(client)
    url = lambda kind, cid=comp_id: f"/api/exports/{kind}/{cid}.pdf"

    # chief judge gets both documents, the secretary only the panel
    assert client.get(url("panel"), headers=headers["judge"]).status_code == 200
    assert client.get(url("chief-report"), headers=headers["judge"]).status_code == 200
    assert client.get(url("panel"), headers=as_judge_admin).status_code == 200
    assert client.get(url("chief-report"), headers=as_judge_admin).status_code == 403
    # the ministry reads, athletes and coaches do not
    assert client.get(url("panel"), headers=ministry).status_code == 200
    assert client.get(url("chief-report"), headers=ministry).status_code == 200
    for role in ("athlete", "coach"):
        assert client.get(url("panel"), headers=headers[role]).status_code == 403
        assert client.get(url("chief-report"), headers=headers[role]).status_code == 403

    with factory() as db:
        chief = db.scalar(select(Judge).where(Judge.user_id == users["judge"]["id"]))
        template = db.get(Competition, comp_id)

        def copy(state, judged):
            c = Competition(
                **{
                    k: getattr(template, k)
                    for k in Competition.__table__.columns.keys()
                    if k != "id"
                },
            )
            c.title, c.state = "Тест " + state, state
            db.add(c)
            db.flush()
            if judged:
                db.add(
                    CompetitionJudge(
                        competition_id=c.id, judge_id=chief.id, position="Главный судья"
                    )
                )
            return c.id

        other, unpublished, draft = (
            copy("closed", False),
            copy("closed", True),
            copy("draft", True),
        )
        db.commit()

    # a judge outside the panel; a panel that is empty
    assert client.get(url("panel", other), headers=headers["judge"]).status_code == 403
    assert client.get(url("panel", other), headers=headers["admin"]).status_code == 409
    # the report needs a published protocol
    res = client.get(url("chief-report", unpublished), headers=headers["judge"])
    assert res.status_code == 409
    # drafts are the federation's own business
    assert client.get(url("panel", draft), headers=ministry).status_code == 404
    assert client.get(url("panel", draft), headers=headers["judge"]).status_code == 404
    assert client.get(url("panel", draft), headers=headers["admin"]).status_code == 200
