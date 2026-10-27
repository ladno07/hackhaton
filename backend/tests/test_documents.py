from datetime import date
from io import BytesIO
from docx import Document
from app.config import settings


def docx_text(content):
    doc = Document(BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts += [cell.text for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def published(client, h):
    comps = client.get("/api/competitions", headers=h["admin"]).json()
    return next(c for c in comps if c["state"] == "published")


def test_protocol_has_letterhead_and_signatures(env):
    client, _, h, _ = env
    comp = published(client, h)
    res = client.get(f"/api/exports/protocol/{comp['id']}.docx", headers=h["admin"])
    assert res.status_code == 200, res.text
    content = docx_text(res.content)
    assert settings.fsp_name in content and settings.fsp_requisites in content
    assert "ИТОГОВЫЙ ПРОТОКОЛ" in content and comp["title"] in content
    assert "Главный судья" in content and "Главный секретарь" in content
    assert "(расшифровка подписи)" in content
    pdf = client.get(f"/api/exports/protocol/{comp['id']}.pdf", headers=h["admin"])
    assert pdf.content.startswith(b"%PDF")


def test_team_order_is_draft_with_appendix(env):
    client, _, h, _ = env
    res = client.get("/api/exports/team.docx?order=true", headers=h["admin"])
    assert res.status_code == 200, res.text
    content = docx_text(res.content)
    assert "ПРИКАЗ" in content and "ПРОЕКТ" in content
    assert "ПРИКАЗЫВАЮ:" in content
    # number and date are filled in by hand after approval
    assert "№ ________" in content and "«___»" in content
    assert "Приложение" in content and "№ п/п" in content
    assert settings.fsp_head_position in content


def test_reports_are_dated_and_signed(env):
    client, _, h, _ = env
    today = f"{date.today():%d.%m.%Y}"
    population = client.get("/api/exports/population.docx", headers=h["admin"])
    assert population.status_code == 200, population.text
    content = docx_text(population.content)
    assert "ОТЧЁТ" in content and today in content
    assert "Исполнитель" in content and "Методика подсчёта" in content
    judging = docx_text(
        client.get("/api/exports/judging.docx", headers=h["judge"]).content
    )
    assert "Спортивный судья" in judging and today in judging
