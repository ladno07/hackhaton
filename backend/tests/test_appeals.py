import io
from datetime import date, timedelta
from openpyxl import load_workbook
from sqlalchemy import select
from app.models import (
    Appeal,
    Athlete,
    AthleteCoach,
    Competition,
    CompetitionResult,
    ProtocolVersion,
    Violation,
    now,
)
from tests.test_judging_documents import docx_text
from tests.test_ministry import ministry_login

CUP = "Демо · Кубок регионов по алгоритмам"
REASON = "Решение задачи F не засчитано, хотя прошло все тесты"


def cup(factory, days_ago=0):
    """Demo cup id; every demo protocol is dated `days_ago`, so only the cup
    stays open for appeals by default and none with days_ago > 3."""
    with factory() as db:
        comp = db.scalar(select(Competition).where(Competition.title == CUP))
        for v in db.scalars(select(ProtocolVersion)):
            fresh = v.competition_id == comp.id and not days_ago
            v.created_at = now() - timedelta(days=0 if fresh else days_ago or 30)
        db.commit()
        return comp.id


def athlete_id(factory, user_id):
    with factory() as db:
        return db.scalar(select(Athlete.id).where(Athlete.user_id == user_id))


def result_of(factory, comp_id, athlete):
    with factory() as db:
        return db.scalar(
            select(CompetitionResult).where(
                CompetitionResult.competition_id == comp_id,
                CompetitionResult.athlete_id == athlete,
            )
        )


def coached(factory, users):
    """A coach's athlete without a demo appeal on the cup."""
    with factory() as db:
        busy = set(db.scalars(select(Appeal.athlete_id)))
        return next(
            l.athlete_id
            for l in db.scalars(select(AthleteCoach).where(AthleteCoach.status == "active"))
            if l.athlete_id not in busy
        )


def titles(client, h):
    return [n["title"] for n in client.get("/api/notifications", headers=h).json()]


def test_result_appeal_upheld_creates_protocol_version(env):
    client, factory, headers, users = env
    comp_id = cup(factory)
    me = athlete_id(factory, users["athlete"]["id"])
    mine = result_of(factory, comp_id, me)

    options = client.get("/api/appeals/options", headers=headers["athlete"]).json()
    assert [o["result_id"] for o in options["results"]] == [mine.id]
    assert options["result_days"] == 3
    # the demo warning was recorded yesterday and may still be appealed
    assert [v["kind"] for v in options["violations"]] == ["conduct"]

    body = {"subject": "result", "result_id": mine.id, "reason": REASON, "request": "Пересчитать баллы"}
    assert client.post("/api/appeals", json={**body, "reason": "коротко"}, headers=headers["athlete"]).status_code == 422
    # the coach has no active link with this athlete
    assert client.post("/api/appeals", json=body, headers=headers["coach"]).status_code == 403
    res = client.post("/api/appeals", json=body, headers=headers["athlete"])
    assert res.status_code == 201, res.text
    appeal = res.json()
    assert appeal["status"] == "pending" and appeal["protocol_version"] == 1
    assert appeal["number"] == f"А-{appeal['id']}" and appeal["can_withdraw"]
    assert "1 место" in appeal["subject_text"]
    assert client.post("/api/appeals", json=body, headers=headers["athlete"]).status_code == 409

    # federation: dashboard and notification; the chief judge of the cup reads it
    attention = client.get("/api/dashboard/attention", headers=headers["admin"]).json()
    assert appeal["id"] in {i["id"] for i in attention["items"] if i["kind"] == "appeals"}
    assert "Новая апелляция" in titles(client, headers["admin"])
    judged = client.get("/api/appeals", headers=headers["judge"]).json()
    assert appeal["id"] in {a["id"] for a in judged}
    assert not any(a["can_decide"] for a in judged)
    assert client.patch(
        f"/api/appeals/{appeal['id']}", json={"status": "rejected", "decision": "Отказ"},
        headers=headers["judge"],
    ).status_code == 403

    detail = client.get(f"/api/competitions/{comp_id}", headers=headers["admin"]).json()
    results = [
        {"athlete_id": r["athlete_id"], "place": r["place"],
         "score": r["score"] + (50 if r["athlete_id"] == me else 0)}
        for r in detail["results"]
    ]
    # a correction is only allowed together with an upheld result appeal
    res = client.patch(
        f"/api/appeals/{appeal['id']}",
        json={"status": "rejected", "decision": "Ошибки нет", "correction": {"results": results}},
        headers=headers["admin"],
    )
    assert res.status_code == 422
    res = client.patch(
        f"/api/appeals/{appeal['id']}",
        json={"status": "upheld", "decision": "Задача F засчитана", "correction": {"results": results}},
        headers=headers["admin"],
    )
    assert res.status_code == 200, res.text
    decided = res.json()
    assert decided["status"] == "upheld" and decided["resolution_version"] == 2
    assert decided["current_score"] == mine.score + 50

    versions = client.get(f"/api/competitions/{comp_id}/protocol-versions", headers=headers["athlete"]).json()
    assert versions[0]["version"] == 2
    assert versions[0]["reason"].startswith(f"Апелляция № А-{appeal['id']}: Задача F засчитана")
    assert "Решение по апелляции" in titles(client, headers["athlete"])
    assert client.patch(
        f"/api/appeals/{appeal['id']}", json={"status": "rejected", "decision": "Повторно"},
        headers=headers["admin"],
    ).status_code == 409

    # the chief judge report lists the appeal
    report = docx_text(
        client.get(f"/api/exports/chief-report/{comp_id}.docx", headers=headers["judge"]).content
    )
    assert "Апелляции на протокол" in report and f"А-{appeal['id']}" in report


def test_deadline_withdraw_and_access(env):
    client, factory, headers, users = env
    comp_id = cup(factory, days_ago=30)
    me = athlete_id(factory, users["athlete"]["id"])
    mine = result_of(factory, comp_id, me)
    body = {"subject": "result", "result_id": mine.id, "reason": REASON}

    # the demo cup was published 30 days ago: the 3-day window has passed
    assert client.get("/api/appeals/options", headers=headers["athlete"]).json()["results"] == []
    res = client.post("/api/appeals", json=body, headers=headers["athlete"])
    assert res.status_code == 409 and "Срок подачи" in res.json()["detail"]

    cup(factory)
    pupil = coached(factory, users)
    theirs = result_of(factory, comp_id, pupil)
    res = client.post(
        "/api/appeals", json={**body, "result_id": theirs.id}, headers=headers["coach"]
    )
    assert res.status_code == 201, res.text
    appeal = res.json()
    assert appeal["submitted_role"] == "coach"
    # the athlete does not see another athlete's appeal and cannot withdraw it
    assert appeal["id"] not in {a["id"] for a in client.get("/api/appeals", headers=headers["athlete"]).json()}
    assert client.post(f"/api/appeals/{appeal['id']}/withdraw", headers=headers["athlete"]).status_code == 404
    res = client.post(f"/api/appeals/{appeal['id']}/withdraw", headers=headers["coach"])
    assert res.status_code == 200 and res.json()["status"] == "withdrawn"
    assert client.post(f"/api/appeals/{appeal['id']}/withdraw", headers=headers["coach"]).status_code == 409

    # the ministry reads and exports, but cannot decide
    ministry = ministry_login(client)
    listed = client.get("/api/appeals", headers=ministry).json()
    assert {"pending", "rejected", "withdrawn"} <= {a["status"] for a in listed}
    pending = next(a for a in listed if a["status"] == "pending")
    assert client.patch(
        f"/api/appeals/{pending['id']}", json={"status": "rejected", "decision": "Отказ"},
        headers=ministry,
    ).status_code == 403
    assert client.get("/api/appeals?status=pending", headers=ministry).json() == [pending]
    sheet = load_workbook(
        io.BytesIO(client.get("/api/exports/appeals.xlsx", headers=ministry).content)
    ).active
    cells = " ".join(str(c.value) for r in sheet.iter_rows() for c in r if c.value)
    assert "апелляций" in cells.lower() and pending["number"] in cells
    assert client.get("/api/exports/appeals.pdf", headers=headers["athlete"]).status_code == 403


def test_violation_suspension_and_appeal(env):
    client, factory, headers, users = env
    pupil = coached(factory, users)
    comp_id = cup(factory, days_ago=30)
    month, two = date.today() + timedelta(days=30), date.today() + timedelta(days=60)
    body = {
        "athlete_id": pupil,
        "competition_id": comp_id,
        "kind": "conduct",
        "occurred_on": str(date.today()),
        "description": "Подсказки участнику во время тура",
        "sanction": "suspension",
        "suspended_until": str(month),
    }
    assert client.post("/api/violations", json={**body, "suspended_until": None}, headers=headers["admin"]).status_code == 422
    assert client.post("/api/violations", json={**body, "sanction": "warning"}, headers=headers["admin"]).status_code == 422
    assert client.post("/api/violations", json=body, headers=headers["coach"]).status_code == 403
    first = client.post("/api/violations", json=body, headers=headers["admin"])
    assert first.status_code == 201, first.text
    second = client.post(
        "/api/violations", json={**body, "suspended_until": str(two)}, headers=headers["admin"]
    ).json()
    with factory() as db:
        assert db.get(Athlete, pupil).suspended_until == two

    # the coach sees the violations of their athlete and appeals the longer one
    seen = client.get("/api/violations", headers=headers["coach"]).json()
    assert {first.json()["id"], second["id"]} <= {v["id"] for v in seen}
    assert all(v["can_appeal"] for v in seen if v["athlete_id"] == pupil)
    assert pupil not in {v["athlete_id"] for v in client.get("/api/violations", headers=headers["athlete"]).json()}
    res = client.post(
        "/api/appeals",
        json={"subject": "violation", "violation_id": second["id"], "reason": "Спортсмен не получал подсказок, есть видеозапись"},
        headers=headers["coach"],
    )
    assert res.status_code == 201, res.text
    appeal = res.json()
    assert appeal["subject_text"].startswith("Решение о нарушении: Неспортивное поведение")

    res = client.patch(
        f"/api/appeals/{appeal['id']}",
        json={"status": "upheld", "decision": "Видеозапись подтверждает доводы",
              "correction": {"results": [{"athlete_id": pupil, "place": 1, "score": 1}]}},
        headers=headers["admin"],
    )
    assert res.status_code == 422
    res = client.patch(
        f"/api/appeals/{appeal['id']}",
        json={"status": "upheld", "decision": "Видеозапись подтверждает доводы"},
        headers=headers["admin"],
    )
    assert res.status_code == 200, res.text
    with factory() as db:
        violation = db.get(Violation, second["id"])
        assert violation.status == "cancelled"
        assert violation.cancel_reason.startswith(f"Апелляция № А-{appeal['id']} удовлетворена")
        # the remaining suspension takes over
        assert db.get(Athlete, pupil).suspended_until == month

    res = client.post(
        f"/api/violations/{first.json()['id']}/cancel", json={"reason": "Ошибка при вводе"},
        headers=headers["admin"],
    )
    assert res.status_code == 200 and res.json()["status"] == "cancelled"
    assert client.post(
        f"/api/violations/{first.json()['id']}/cancel", json={"reason": "Повторно"},
        headers=headers["admin"],
    ).status_code == 409
    with factory() as db:
        assert db.get(Athlete, pupil).suspended_until is None


def test_manual_suspension_is_kept(env):
    client, factory, headers, users = env
    pupil = coached(factory, users)
    manual = date.today() + timedelta(days=90)
    res = client.put(
        f"/api/athletes/{pupil}/suspension",
        json={"suspended_until": str(manual), "reason": "Решение дисциплинарной комиссии"},
        headers=headers["admin"],
    )
    assert res.status_code == 200, res.text
    violation = client.post(
        "/api/violations",
        json={
            "athlete_id": pupil,
            "kind": "other",
            "occurred_on": str(date.today()),
            "description": "Нарушение регламента",
            "sanction": "suspension",
            "suspended_until": str(date.today() + timedelta(days=10)),
        },
        headers=headers["admin"],
    ).json()
    client.post(f"/api/violations/{violation['id']}/cancel", json={"reason": "Ошибка"}, headers=headers["admin"])
    with factory() as db:
        athlete = db.get(Athlete, pupil)
        assert athlete.suspended_until == manual
        assert athlete.suspension_reason == "Решение дисциплинарной комиссии"
