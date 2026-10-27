"""Official document template: letterhead with FSP requisites, document kind,
date and number line, body, table and signature block. Rendered to PDF, DOCX
and XLSX from one description, so every export looks like the same form."""

from dataclasses import dataclass, field
from datetime import date
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from fastapi import HTTPException
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from app.config import settings

MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
BLANK_DATE = "«___» ____________ 20___ г."
SIGN_LINE = "_________________"
DEMO_NOTE = "Демонстрационные данные"


@dataclass
class OfficialDocument:
    kind: str  # "Итоговый протокол", "Приказ", "Отчёт"...
    title: str  # lowercase continues the kind: "о численности ..."
    headers: list
    rows: list
    meta: list = field(default_factory=list)  # (label, value) pairs
    intro: list = field(default_factory=list)  # paragraphs before the table
    outro: list = field(default_factory=list)  # paragraphs after the table
    signatures: list = field(default_factory=list)  # (position, full name)
    approve: tuple | None = None  # "УТВЕРЖДАЮ" block: (position, full name)
    dated: date | None = None  # None leaves a blank to fill in by hand
    number: str | None = None  # None: no number; "": blank to fill in
    place: str = ""
    draft: bool = False
    appendix: list = field(default_factory=list)  # table goes to an appendix
    table_title: str = ""

    @property
    def heading(self):
        joint = " " if self.title[:1].islower() else ". "
        return ("ПРОЕКТ. " if self.draft else "") + self.kind + joint + self.title

    @property
    def date_line(self):
        return f"{self.dated:%d.%m.%Y}" if self.dated else BLANK_DATE

    @property
    def number_line(self):
        if self.number is None:
            return ""
        return "№ " + (self.number or "________")


def signature_name(full_name):
    # Names are stored free-form ("Имя Фамилия"), so they are not reduced to
    # initials: a wrong guess would be worse than the full name.
    return full_name or SIGN_LINE


def safe_cell(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def text(value):
    return "" if value is None else str(value)


def column_weights(doc, capacity=95):
    """Relative column widths. Every column first gets room for its longest
    word, so nothing breaks mid-word; the rest of the line (capacity, in
    characters) goes to columns in proportion to their average text."""
    need, flex = [], []
    for i, header in enumerate(doc.headers):
        cells = [text(r[i]) for r in doc.rows] or [""]
        words = [len(w) for c in cells for w in c.split()]
        # header is bold, so its letters are a bit wider
        words += [len(w) * 1.1 for w in header.split()]
        need.append(max(words) + 2)
        flex.append(sum(len(c) for c in cells) / len(cells) + 1)
    spare = capacity - sum(need)
    if spare > 0:
        need = [n + spare * f / sum(flex) for n, f in zip(need, flex)]
    total = sum(need)
    return [n / total for n in need]


def export_document(doc: OfficialDocument, fmt):
    render = {"pdf": render_pdf, "docx": render_docx, "xlsx": render_xlsx}.get(fmt)
    if not render:
        raise HTTPException(422, "Формат: pdf, docx или xlsx")
    doc.place = doc.place or settings.fsp_city
    out = BytesIO()
    render(doc, out)
    return out.getvalue()


# --- XLSX: data first, so the table keeps fixed rows for filtering ---------


def render_xlsx(doc, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws.append([safe_cell(doc.heading)])
    ws.append(
        [safe_cell("; ".join(f"{label}: {value}" for label, value in doc.meta))]
    )
    ws.append(doc.headers)
    for values in doc.rows:
        ws.append([safe_cell(v) for v in values])
    last = ws.max_row
    ws["A1"].font = Font(bold=True, size=13)
    thin = Side(style="thin", color="999999")
    for row in ws.iter_rows(min_row=3, max_row=last, max_col=len(doc.headers)):
        for c in row:
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[3]:
        c.font = Font(color="FFFFFF", bold=True)
        c.fill = PatternFill("solid", fgColor="B62C3B")
    for column in ws.iter_cols(min_row=3, max_row=last):
        width = min(55, max(14, max(len(text(c.value)) for c in column) + 2))
        ws.column_dimensions[column[0].column_letter].width = width
        for c in column:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.append([])
    for paragraph in doc.outro:
        ws.append([safe_cell(paragraph)])
    for position, name in doc.signatures:
        ws.append([position, SIGN_LINE, signature_name(name)])
    ws.append([f"{settings.fsp_short_name} · {doc.place} · {doc.date_line}"])
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{ws.cell(max(3, last), len(doc.headers)).coordinate}"
    ws.oddHeader.center.text = settings.fsp_name
    ws.oddHeader.center.size = 8
    ws.oddFooter.left.text = DEMO_NOTE
    ws.oddFooter.right.text = "Стр. &P из &N"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "3:3"
    wb.save(out)


# --- PDF ---------------------------------------------------------------------


def pdf_fonts():
    regular = Path(settings.pdf_font)
    if not regular.is_file():
        raise HTTPException(503, "Не найден PDF_FONT с поддержкой кириллицы")
    if "Doc" not in pdfmetrics.getRegisteredFontNames():
        # DejaVuSans.ttf -> DejaVuSans-Bold.ttf, arial.ttf -> arialbd.ttf
        bold = next(
            (
                p
                for p in (
                    regular.with_name(regular.stem + "-Bold" + regular.suffix),
                    regular.with_name(regular.stem + "bd" + regular.suffix),
                )
                if p.is_file()
            ),
            regular,
        )
        pdfmetrics.registerFont(TTFont("Doc", str(regular)))
        pdfmetrics.registerFont(TTFont("Doc-Bold", str(bold)))
        pdfmetrics.registerFontFamily(
            "Doc", normal="Doc", bold="Doc-Bold", italic="Doc", boldItalic="Doc-Bold"
        )


def render_pdf(doc, out):
    pdf_fonts()
    page_w, page_h = A4
    left, right = 25 * mm, 15 * mm
    width = page_w - left - right
    base = ParagraphStyle("base", fontName="Doc", fontSize=10.5, leading=14)
    s = {
        "org": ParagraphStyle(
            "org", parent=base, fontName="Doc-Bold", fontSize=10, leading=13,
            alignment=TA_CENTER,
        ),
        "req": ParagraphStyle(
            "req", parent=base, fontSize=7.5, leading=10, alignment=TA_CENTER,
            textColor=colors.HexColor("#444444"),
        ),
        "kind": ParagraphStyle(
            "kind", parent=base, fontName="Doc-Bold", fontSize=14, leading=18,
            alignment=TA_CENTER, spaceBefore=6,
        ),
        "title": ParagraphStyle(
            "title", parent=base, fontName="Doc-Bold", fontSize=11.5, leading=15,
            alignment=TA_CENTER, spaceAfter=10,
        ),
        "meta": ParagraphStyle("meta", parent=base, fontSize=9.5, leading=13),
        "body": ParagraphStyle(
            "body", parent=base, alignment=TA_JUSTIFY, firstLineIndent=12.5 * mm,
            spaceAfter=4,
        ),
        "order": ParagraphStyle(
            "order", parent=base, fontName="Doc-Bold", spaceBefore=4, spaceAfter=4
        ),
        "right": ParagraphStyle("right", parent=base, fontSize=9.5, leading=12,
                                alignment=TA_RIGHT),
        "cell": ParagraphStyle("cell", parent=base, fontSize=8.5, leading=11),
        "head": ParagraphStyle(
            "head", parent=base, fontName="Doc-Bold", fontSize=8.5, leading=11,
            alignment=TA_CENTER,
        ),
        "sign": ParagraphStyle("sign", parent=base, fontSize=10.5, leading=13),
        "hint": ParagraphStyle(
            "hint", parent=base, fontSize=6.5, leading=8, alignment=TA_CENTER,
            textColor=colors.HexColor("#666666"),
        ),
    }
    p = lambda value, style: Paragraph(
        escape(text(value)).replace("\n", "<br/>"), s[style]
    )
    plain = TableStyle(
        [
            ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ]
    )

    story = [
        p(settings.fsp_name, "org"),
        Spacer(1, 2),
        p(settings.fsp_requisites, "req"),
        HRFlowable(width="100%", thickness=1.4, color=colors.black, spaceBefore=4,
                   spaceAfter=1),
        HRFlowable(width="100%", thickness=0.5, color=colors.black, spaceAfter=8),
    ]
    if doc.draft:
        story.append(Paragraph("<b>ПРОЕКТ</b>", s["right"]))
    if doc.approve:
        position, name = doc.approve
        block = Table(
            [
                ["", p("УТВЕРЖДАЮ", "order")],
                ["", p(f"{position}\n{settings.fsp_short_name}", "sign")],
                ["", p(f"{SIGN_LINE} {signature_name(name)}", "sign")],
                ["", p(BLANK_DATE, "sign")],
            ],
            colWidths=[width * 0.5, width * 0.5],
        )
        block.setStyle(plain)
        story += [block, Spacer(1, 10)]
    story.append(p(doc.kind.upper(), "kind"))
    line = Table(
        [[p(doc.date_line, "sign"), p(doc.place, "head"), p(doc.number_line, "right")]],
        colWidths=[width * 0.35, width * 0.3, width * 0.35],
    )
    line.setStyle(plain)
    story += [Spacer(1, 4), line, Spacer(1, 8), p(doc.title, "title")]
    for label, value in doc.meta:
        story.append(
            Paragraph(f"<b>{escape(label)}:</b> {escape(text(value))}", s["meta"])
        )
    if doc.meta:
        story.append(Spacer(1, 8))
    for paragraph in doc.intro:
        story.append(p(paragraph, "order" if paragraph.isupper() else "body"))

    weights = column_weights(doc)
    table = Table(
        [[p(h, "head") for h in doc.headers]]
        + [[p(v, "cell") for v in r] for r in doc.rows],
        repeatRows=1,
        colWidths=[width * w for w in weights],
        hAlign="LEFT",
    )
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6E6E6")),
                ("VALIGN", (0, 0), (-1, 0), "MIDDLE"),
                ("VALIGN", (0, 1), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    signatures = [Spacer(1, 22)]
    for position, name in doc.signatures:
        row = Table(
            [
                [p(position, "sign"), p(SIGN_LINE, "sign"), p(signature_name(name), "sign")],
                ["", p("(подпись)", "hint"), p("(расшифровка подписи)", "hint")],
            ],
            colWidths=[width * 0.42, width * 0.28, width * 0.3],
        )
        row.setStyle(plain)
        signatures += [row, Spacer(1, 12)]
    signatures = [KeepTogether(signatures)] if doc.signatures else []

    body = [table]
    if doc.table_title:
        body.insert(0, p(doc.table_title, "title"))
    body += [Spacer(1, 8)] + [p(x, "body") for x in doc.outro]
    if doc.appendix:
        story += signatures + [PageBreak(), p("\n".join(doc.appendix), "right"),
                               Spacer(1, 12)] + body
    else:
        story += body + signatures

    generated = f"{date.today():%d.%m.%Y}"

    def footer(canvas, pdf):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#999999"))
        canvas.setLineWidth(0.4)
        canvas.line(left, 14 * mm, page_w - right, 14 * mm)
        canvas.setFont("Doc", 7)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(
            left, 10 * mm,
            f"{settings.fsp_short_name} · {doc.kind} · сформировано {generated} · {DEMO_NOTE}",
        )
        canvas.drawRightString(page_w - right, 10 * mm, f"Стр. {pdf.page}")
        canvas.restoreState()

    SimpleDocTemplate(
        out,
        pagesize=A4,
        leftMargin=left,
        rightMargin=right,
        topMargin=18 * mm,
        bottomMargin=20 * mm,
        title=doc.heading,
        author=settings.fsp_short_name,
    ).build(story, onFirstPage=footer, onLaterPages=footer)


# --- DOCX --------------------------------------------------------------------


def docx_paragraph(target, value="", size=12, bold=False, align=None, after=0,
                   indent=None, color=None):
    paragraph = target.add_paragraph()
    run = paragraph.add_run(text(value))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(after)
    fmt.space_before = Pt(0)
    if indent is not None:
        fmt.first_line_indent = indent
    if align is not None:
        paragraph.alignment = align
    return paragraph


def docx_rule(paragraph):
    """Double line under the letterhead, as on a federation blank."""
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in {"val": "thinThickSmallGap", "sz": "12", "space": "4",
                       "color": "000000"}.items():
        bottom.set(qn("w:" + key), value)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().insert_element_before(
        border, "w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc"
    )


def docx_layout(doc, widths):
    """Borderless table used to place blocks side by side."""
    table = doc.add_table(rows=1, cols=len(widths))
    table.autofit = False
    for cell, w in zip(table.rows[0].cells, widths):
        cell.width = w
    return table


def docx_cell(cell, value, size=12, bold=False, align=None, color=None):
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text(value))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    return paragraph


def docx_page_number(paragraph):
    for kind, instruction in (("begin", None), (None, " PAGE "), ("end", None)):
        run = paragraph.add_run()
        run.font.size = Pt(8)
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = instruction
        run._r.append(element)


def render_docx(doc, out):
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0)
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin, section.right_margin = Mm(25), Mm(15)
    section.top_margin, section.bottom_margin = Mm(18), Mm(20)
    width = Mm(170)
    center, right = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer.paragraphs[0]
    footer.alignment = right
    run = footer.add_run(
        f"{settings.fsp_short_name} · сформировано {date.today():%d.%m.%Y} · "
        f"{DEMO_NOTE} · стр. "
    )
    run.font.size = Pt(8)
    docx_page_number(footer)

    docx_paragraph(document, settings.fsp_name, 11, True, center)
    requisites = document.add_paragraph()
    docx_rule(requisites)
    run = requisites.add_run(settings.fsp_requisites)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("444444")
    requisites.alignment = center
    requisites.paragraph_format.space_after = Pt(10)

    if doc.draft:
        docx_paragraph(document, "ПРОЕКТ", 11, True, right, after=4)
    if doc.approve:
        position, name = doc.approve
        cell = docx_layout(document, [width // 2, width // 2]).rows[0].cells[1]
        docx_cell(cell, "УТВЕРЖДАЮ", bold=True)
        for line in (position, settings.fsp_short_name,
                     f"{SIGN_LINE} {signature_name(name)}", BLANK_DATE):
            docx_paragraph(cell, line)
        docx_paragraph(document, after=6)

    docx_paragraph(document, doc.kind.upper(), 14, True, center, after=6)
    cells = docx_layout(document, [width * 35 // 100, width * 30 // 100,
                                   width * 35 // 100]).rows[0].cells
    docx_cell(cells[0], doc.date_line)
    docx_cell(cells[1], doc.place, align=center)
    docx_cell(cells[2], doc.number_line, align=right)
    docx_paragraph(document, after=6)
    docx_paragraph(document, doc.title, 12, True, center, after=10)
    for label, value in doc.meta:
        paragraph = docx_paragraph(document, label + ": ", 11, True)
        paragraph.add_run(text(value)).font.size = Pt(11)
    if doc.meta:
        docx_paragraph(document, after=4)
    for paragraph in doc.intro:
        if paragraph.isupper():
            docx_paragraph(document, paragraph, bold=True, after=4)
        else:
            docx_paragraph(document, paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                           after=4, indent=Mm(12.5))

    def add_signatures():
        docx_paragraph(document, after=18)
        for position, name in doc.signatures:
            row = docx_layout(document, [width * 42 // 100, width * 28 // 100,
                                         width * 30 // 100]).rows[0].cells
            docx_cell(row[0], position)
            docx_cell(row[1], SIGN_LINE)
            docx_cell(row[2], signature_name(name))
            hint = docx_layout(document, [width * 42 // 100, width * 28 // 100,
                                          width * 30 // 100]).rows[0].cells
            docx_cell(hint[1], "(подпись)", 7, align=center, color="666666")
            docx_cell(hint[2], "(расшифровка подписи)", 7, align=center,
                      color="666666")
            docx_paragraph(document, after=8)

    if doc.appendix:
        add_signatures()
        page = docx_paragraph(document)
        page.runs[0].add_break(WD_BREAK.PAGE)
        for line in doc.appendix:
            docx_paragraph(document, line, 11, align=right)
        docx_paragraph(document, after=8)
    if doc.table_title:
        docx_paragraph(document, doc.table_title, 12, True, center, after=8)

    weights = column_weights(doc)
    table = document.add_table(rows=1, cols=len(doc.headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    header._tr.get_or_add_trPr().append(flag)
    for cell, h in zip(header.cells, doc.headers):
        docx_cell(cell, h, 10, True, center)
        shade = OxmlElement("w:shd")
        for key, value in {"val": "clear", "color": "auto", "fill": "E6E6E6"}.items():
            shade.set(qn("w:" + key), value)
        cell._tc.get_or_add_tcPr().append(shade)
    for values in doc.rows:
        for cell, v in zip(table.add_row().cells, values):
            docx_cell(cell, v, 10)
    for row in table.rows:
        for cell, w in zip(row.cells, weights):
            cell.width = int(width * w)
    for paragraph in doc.outro:
        docx_paragraph(document, paragraph, 11, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                       indent=Mm(12.5)).paragraph_format.space_before = Pt(6)
    if not doc.appendix:
        add_signatures()
    document.core_properties.title = doc.heading
    document.core_properties.author = settings.fsp_short_name
    document.save(out)
