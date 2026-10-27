from collections import defaultdict
from datetime import date
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from sqlalchemy import select
from fastapi import HTTPException
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from app.config import settings
from app.models import *
from app.services.common import age_at

MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def population(
    db,
    start_year,
    end_year,
    organization_id=None,
    discipline_id=None,
    municipality="",
    min_age=0,
    max_age=100,
):
    counts = defaultdict(set)
    for enrollment, athlete in db.execute(
        select(AthleteDiscipline, Athlete).join(
            Athlete, Athlete.id == AthleteDiscipline.athlete_id
        )
    ):
        if (
            organization_id
            and athlete.organization_id != organization_id
            or discipline_id
            and enrollment.discipline_id != discipline_id
            or municipality.lower() not in athlete.municipality.lower()
        ):
            continue
        org = (
            db.get(Organization, athlete.organization_id)
            if athlete.organization_id
            else None
        )
        for year in range(start_year, end_year + 1):
            if not enrollment.start_year <= year <= (enrollment.end_year or end_year):
                continue
            age = age_at(athlete.user.birth_date, date(year, 12, 31))
            if not min_age <= age <= max_age:
                continue
            group = (
                "До 14 лет"
                if age < 14
                else "14–17 лет"
                if age < 18
                else "18 лет и старше"
            )
            key = (
                year,
                org.name if org else "Без организации",
                db.get(Discipline, enrollment.discipline_id).short_name,
                group,
                athlete.municipality,
            )
            counts[key].add(athlete.id)
    rows = [[*key, len(ids)] for key, ids in sorted(counts.items())]
    totals = []
    for year in range(start_year, end_year + 1):
        unique = set().union(*(ids for key, ids in counts.items() if key[0] == year))
        totals.append({"year": year, "count": len(unique)})
    return {
        "totals": totals,
        "rows": rows,
        "headers": [
            "Год",
            "Организация",
            "Дисциплина",
            "Возрастная группа",
            "Муниципалитет",
            "Численность",
        ],
        "source": "Годы начала и окончания занятий из реестра. Возраст на 31 декабря; в итогах спортсмен считается один раз.",
    }


def safe_cell(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def export_document(title, headers, rows, fmt, subtitle=""):
    out = BytesIO()
    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Данные"
        ws.append([safe_cell(title)])
        ws.append([safe_cell(subtitle)])
        ws.append(headers)
        for values in rows:
            ws.append([safe_cell(v) for v in values])
        for c in ws[3]:
            c.font = Font(color="FFFFFF", bold=True)
            c.fill = PatternFill("solid", fgColor="B62C3B")
        for column in ws.columns:
            width = min(
                55, max(14, max(len(str(c.value or "")) for c in list(column)[2:]) + 2)
            )
            ws.column_dimensions[column[0].column_letter].width = width
            for c in column:
                if isinstance(c.value, date):
                    c.number_format = "dd.mm.yyyy"
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = (
            f"A3:{ws.cell(max(3, ws.max_row), len(headers)).coordinate}"
        )
        wb.save(out)
    elif fmt == "docx":
        doc = Document()
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(10)
        doc.add_heading(title, 0)
        doc.add_paragraph(subtitle)
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Shading Accent 1"
        for c, h in zip(table.rows[0].cells, headers):
            c.text = str(h)
        for values in rows:
            for c, v in zip(table.add_row().cells, values):
                c.text = v.strftime("%d.%m.%Y") if isinstance(v, date) else str(v if v is not None else "")
        doc.add_paragraph(
            "Федерация спортивного программирования Оренбургской области. Демонстрационные данные."
        )
        doc.save(out)
    elif fmt == "pdf":
        font = Path(settings.pdf_font)
        if not font.is_file():
            raise HTTPException(503, "Не найден PDF_FONT с поддержкой кириллицы")
        if "Cyrillic" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("Cyrillic", str(font)))
        style = ParagraphStyle("body", fontName="Cyrillic", fontSize=8, leading=12)
        heading = ParagraphStyle(
            "heading", parent=style, fontSize=16, leading=21, spaceAfter=12
        )
        wrap = lambda v: Paragraph(escape(v.strftime("%d.%m.%Y") if isinstance(v, date) else str(v if v is not None else "")), style)
        document = SimpleDocTemplate(
            out,
            pagesize=landscape(A4),
            rightMargin=28,
            leftMargin=28,
            topMargin=30,
            bottomMargin=30,
        )
        table = Table(
            [[wrap(v) for v in headers]] + [[wrap(v) for v in r] for r in rows],
            repeatRows=1,
            colWidths=[(landscape(A4)[0] - 56) / len(headers)] * len(headers),
            hAlign="LEFT",
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3E7E8")),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#F5F6F8")],
                    ),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]
            )
        )

        def footer(canvas, doc):
            canvas.setFont("Cyrillic", 8)
            canvas.drawString(
                28, 16, "ФСП Оренбургской области · Демонстрационные данные"
            )
            canvas.drawRightString(landscape(A4)[0] - 28, 16, str(doc.page))

        document.build(
            [Paragraph(escape(title), heading), wrap(subtitle), Spacer(1, 14), table],
            onFirstPage=footer,
            onLaterPages=footer,
        )
    else:
        raise HTTPException(422, "Формат: pdf, docx или xlsx")
    return out.getvalue()
